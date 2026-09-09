import csv
import io
import re
import zipfile
from pathlib import PurePath
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from app.core.errors import AppError

SUPPORTED = {'.pdf', '.docx', '.pptx', '.txt', '.csv', '.xlsx'}
MAX_CHARS = 500_000


def validate_file(filename, data, max_mb=15):
    name = PurePath(filename.replace('\\', '/')).name[:180]
    ext = PurePath(name).suffix.lower()
    if ext not in SUPPORTED:
        raise AppError('Choose a PDF, DOCX, PPTX, TXT, CSV, or XLSX file.', 415)
    if not data:
        raise AppError('This file is empty.')
    if len(data) > max_mb * 1024 * 1024:
        raise AppError(f'Files must be {max_mb} MB or smaller.', 413)
    if ext == '.pdf' and not data.startswith(b'%PDF-'):
        raise AppError('The file contents do not match a valid PDF.')
    if ext in ('.docx', '.xlsx', '.pptx'):
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise AppError('The file contents do not match an Office document.')
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise AppError('This document expands beyond the safe processing limit.', 413)
            required = {'.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml', '.pptx': 'ppt/presentation.xml'}[ext]
            if required not in archive.namelist():
                raise AppError('The document format does not match its extension.')
    return name, ext


def extract(data: bytes, ext: str):
    pages = []
    stream = io.BytesIO(data)
    try:
        if ext == '.pdf':
            reader = PdfReader(stream)
            if len(reader.pages) > 300:
                raise AppError('Please upload a PDF with 300 pages or fewer.')
            pages = [(n + 1, p.extract_text() or '') for n, p in enumerate(reader.pages)]
        elif ext == '.docx':
            doc = Document(stream)
            text = '\n'.join(p.text for p in doc.paragraphs)
            text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
            pages = [(1, text)]
        elif ext == '.pptx':
            presentation = Presentation(stream)
            if len(presentation.slides) > 300:
                raise AppError('Please upload a presentation with 300 slides or fewer.')
            for i, slide in enumerate(presentation.slides):
                text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text.append(shape.text)
                    if shape.has_table:
                        text.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                if slide.has_notes_slide:
                    text.append(slide.notes_slide.notes_text_frame.text)
                pages.append((i + 1, '\n'.join(text)))
        elif ext == '.xlsx':
            wb = load_workbook(stream, read_only=True, data_only=True)
            try:
                for sheet in wb:
                    rows = []
                    for i, row in enumerate(sheet.iter_rows(values_only=True)):
                        if i >= 10000:
                            raise AppError('Please reduce the workbook to 10,000 rows per sheet or fewer.')
                        rows.append(' | '.join('' if v is None else str(v) for v in row))
                    pages.append((sheet.title, '\n'.join(rows)))
            finally:
                wb.close()
        else:
            text = data.decode('utf-8-sig')
            if '\x00' in text:
                raise AppError('Please upload a UTF-8 text file, not binary data.')
            if ext == '.csv':
                text = '\n'.join(' | '.join(row) for row in csv.reader(io.StringIO(text)))
            pages = [(1, text)]
    except AppError:
        raise
    except Exception:
        raise AppError('This file could not be read. Check that it is not encrypted or damaged. Text files must use UTF-8.') from None
    if sum(len(t) for _, t in pages) > MAX_CHARS:
        raise AppError('This document contains too much text. Split it into smaller files.')
    if not any(t.strip() for _, t in pages):
        raise AppError('No readable text found. Scanned PDFs need OCR before upload.')
    chunks = []
    for page, text in pages:
        text = re.sub(r'[ \t]+', ' ', text).strip()
        for start in range(0, len(text), 1600):
            chunks.append({'text': text[start:start + 1800], 'metadata': {'page': page}, 'chunk_index': len(chunks)})
    return chunks
