"""Generate editable Office documents and BPMN without executing model code."""
import io
import json
import textwrap
import zipfile
from xml.etree.ElementTree import Element, SubElement, tostring
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from app.core.errors import AppError

MIME = {'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        'json': 'application/json', 'md': 'text/markdown; charset=utf-8', 'zip': 'application/zip', 'bpmn': 'application/xml'}


def lines(value, depth=0):
    if isinstance(value, dict) and 'final_report' in value:
        yield from report_lines(value['final_report'])
        return
    if isinstance(value, dict):
        for key, child in value.items():
            yield depth, key.replace('_', ' ').title()
            yield from lines(child, depth + 1)
    elif isinstance(value, list):
        for child in value:
            yield from lines(child, depth)
    elif value is not None:
        yield depth, str(value)


def report_lines(report):
    yield 0, report['title']
    if report.get('report_version'): yield 2, f"Version {report['report_version']} · Prepared {report['prepared_at']}"
    yield 1, f"Industry: {report['industry']} | Language: {report['language']} | Selected option: {report['selected_option']}"
    for i, section in enumerate(report['sections'], 2):
        yield 0, f"{i}. {section['title']}"
        if section['applicability'] == 'not_applicable': yield 2, 'Not applicable — rationale below.'
        yield 2, section['narrative']
        yield 2, 'Evidence / assumptions: ' + section['basis']
        for item in section['items']: yield 2, item
        for t in section['tables']:
            yield 1, t['title']
            yield 2, ' | '.join(t['columns'])
            for row in t['rows']: yield 2, ' | '.join(row)
        yield from visual_lines(section)


def visual_lines(section):
    for diagram in section['diagrams']:
        yield 1, diagram['title']
        names = {n['id']: n['label'] for n in diagram['nodes']}
        for node in diagram['nodes']: yield 2, f"{node['label']} ({node['lane']}; {node['kind']})"
        for edge in diagram['edges']: yield 2, f"{names[edge['source']]} → {names[edge['target']]}: {edge['label']}"
    for screen in section['screens']:
        yield 1, screen['name']
        yield 2, screen['persona'] + ': ' + screen['purpose']
        for control in screen['controls']: yield 2, f"{control['kind']} / {control['label']}: {control['detail']}"
    for asset in section['code_assets']:
        yield 1, asset['filename']
        yield 2, asset['content']


def report_docx(title, report):
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph('shift.AI · Strategy & implementation blueprint')
    if report.get('report_version'): doc.add_paragraph(f"Version {report['report_version']} · Prepared {report['prepared_at']}")
    doc.add_paragraph(f"Industry: {report['industry']}\nLanguage: {report['language']}\nSelected option: {report['selected_option']}")
    doc.add_paragraph('Advisory design. Validate evidence, assumptions and material risks before implementation.')
    doc.add_page_break()
    doc.add_heading('Report contents', 1)
    for i, section in enumerate(report['sections'], 2): doc.add_paragraph(f"{i}. {section['title']}")
    doc.add_page_break()
    for i, section in enumerate(report['sections'], 2):
        doc.add_heading(f"{i}. {section['title']}", 1)
        if section['applicability'] == 'not_applicable': doc.add_paragraph('Not applicable — rationale below.')
        doc.add_paragraph(section['narrative'])
        doc.add_paragraph('Evidence / assumptions: ' + section['basis'])
        for item in section['items']: doc.add_paragraph(item, style='List Bullet')
        for t in section['tables']:
            doc.add_heading(t['title'], 2)
            grid = doc.add_table(rows=1, cols=len(t['columns']))
            grid.style = 'Light Shading Accent 1'
            for cell, value in zip(grid.rows[0].cells, t['columns']): cell.text = value
            for row in t['rows']:
                for cell, value in zip(grid.add_row().cells, row): cell.text = value
        for depth, text in visual_lines(section):
            if depth == 1: doc.add_heading(text, 2)
            else: doc.add_paragraph(text)
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def bpmn(diagram):
    ns = 'http://www.omg.org/spec/BPMN/20100524/MODEL'
    di = 'http://www.omg.org/spec/BPMN/20100524/DI'
    dc = 'http://www.omg.org/spec/DD/20100524/DC'
    dd = 'http://www.omg.org/spec/DD/20100524/DI'
    root = Element('definitions', {'xmlns': ns, 'xmlns:bpmndi': di, 'xmlns:dc': dc, 'xmlns:di': dd, 'targetNamespace': 'https://shift.ai/design', 'id': 'Definitions_1'})
    process = SubElement(root, 'process', {'id': 'Process_1', 'isExecutable': 'false', 'name': diagram['title']})
    lane_set = SubElement(process, 'laneSet', {'id': 'Lanes_1'})
    lanes = list(dict.fromkeys(n['lane'] or 'Process' for n in diagram['nodes']))
    for i, lane in enumerate(lanes):
        element = SubElement(lane_set, 'lane', {'id': f'Lane_{i}', 'name': lane})
        for node in diagram['nodes']:
            if (node['lane'] or 'Process') == lane:
                SubElement(element, 'flowNodeRef').text = node['id']
    types = {'start':'startEvent','end':'endEvent','decision':'exclusiveGateway'}
    positions = {}
    offsets = {lane: 0 for lane in lanes}
    for node in diagram['nodes']:
        tag = types.get(node['kind'], 'task')
        SubElement(process, tag, {'id': node['id'], 'name': node['label']})
        lane = node['lane'] or 'Process'
        positions[node['id']] = (100 + offsets[lane] * 220, 70 + lanes.index(lane) * 150)
        offsets[lane] += 1
    for i, edge in enumerate(diagram['edges']):
        SubElement(process, 'sequenceFlow', {'id': f'Flow_{i}', 'sourceRef': edge['source'], 'targetRef': edge['target'], 'name': edge['label']})
    plane = SubElement(SubElement(root, 'bpmndi:BPMNDiagram', {'id':'Diagram_1'}), 'bpmndi:BPMNPlane', {'id':'Plane_1','bpmnElement':'Process_1'})
    for node in diagram['nodes']:
        x, y = positions[node['id']]
        shape = SubElement(plane, 'bpmndi:BPMNShape', {'id': 'Shape_' + node['id'], 'bpmnElement': node['id']})
        SubElement(shape, 'dc:Bounds', {'x':str(x),'y':str(y),'width':'160','height':'70'})
    for i, edge in enumerate(diagram['edges']):
        shape = SubElement(plane, 'bpmndi:BPMNEdge', {'id':f'Edge_{i}', 'bpmnElement':f'Flow_{i}'})
        x, y = positions[edge['source']]; tx, ty = positions[edge['target']]
        SubElement(shape, 'di:waypoint', {'x':str(x+160),'y':str(y+35)})
        SubElement(shape, 'di:waypoint', {'x':str(tx),'y':str(ty+35)})
    return tostring(root, encoding='utf-8', xml_declaration=True)


def export(title, content, fmt):
    if fmt not in MIME: raise AppError('Unsupported export format.', 415)
    if fmt == 'json': return json.dumps(content, ensure_ascii=False, indent=2).encode()
    report = content.get('final_report')
    if report and fmt == 'docx': return report_docx(title, report)
    sections = report['sections'] if report else content.get('sections', [])
    diagrams = [d for s in sections for d in s.get('diagrams', [])] if report else content.get('diagrams', [])
    assets = [a for s in sections for a in s.get('code_assets', [])] if report else content.get('code_assets', [])
    flattened = list(lines(content))
    if fmt == 'md': return ('# ' + title + '\n\n' + '\n\n'.join('  ' * min(depth, 4) + '- ' + text for depth, text in flattened)).encode()
    if fmt == 'bpmn':
        diagram = next((d for d in diagrams if d['kind'] in ['bpmn','workflow','swimlane']), None)
        if not diagram: raise AppError('Generate a process diagram before exporting BPMN.', 409)
        return bpmn(diagram)
    if fmt == 'zip':
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('deliverable.json', export(title, content, 'json'))
            archive.writestr('report.md', export(title, content, 'md'))
            for asset in assets:
                # Schemas validate filenames, but defend older persisted data too.
                name = asset['filename']
                if '/' not in name and '\\' not in name and '..' not in name:
                    archive.writestr(name, asset['content'])
            for i, diagram in enumerate(diagrams):
                if diagram['kind'] in ['bpmn','workflow','swimlane']:
                    archive.writestr(f'process-{i+1}.bpmn', bpmn(diagram))
        return out.getvalue()
    out = io.BytesIO()
    if fmt == 'docx':
        doc = Document(); doc.add_heading(title, 0)
        doc.add_paragraph('Prepared with shift.AI. Advisory design; validate assumptions before implementation.')
        for depth, text in flattened:
            if depth <= 1 and len(text) < 150: doc.add_heading(text, min(depth + 1, 3))
            else: doc.add_paragraph(text)
        doc.save(out)
    elif fmt == 'xlsx':
        wb = Workbook(); sheet = wb.active; sheet.title = 'Report'; sheet.append(['Level', 'Content'])
        for depth, text in flattened:
            cell = sheet.cell(sheet.max_row + 1, 2, text[:32767]); cell.data_type = 's'
            sheet.cell(cell.row, 1, depth)
        sheet.column_dimensions['B'].width = 100
        for section in sections:
            for table in section.get('tables', []):
                tab = wb.create_sheet('Table ' + str(len(wb.worksheets)))
                for i, row in enumerate([table['columns'], *table['rows']], 1):
                    for j, value in enumerate(row, 1):
                        cell = tab.cell(i, j, str(value)[:32767]); cell.data_type = 's'
        wb.save(out)
    else:
        deck = Presentation(); deck.slide_width = Inches(13.33); deck.slide_height = Inches(7.5)
        chunks = []
        for _, text in flattened:
            chunks.extend(textwrap.wrap(text, 105, replace_whitespace=False) or [''])
        for start in range(0, len(chunks), 13):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            heading = slide.shapes.add_textbox(Inches(.6), Inches(.35), Inches(12), Inches(.8)).text_frame
            heading.text = title[:100] + (f' · {start//13+1}' if start else '')
            heading.paragraphs[0].font.size = Pt(26)
            frame = slide.shapes.add_textbox(Inches(.6), Inches(1.3), Inches(12), Inches(5.7)).text_frame
            frame.word_wrap = True
            for i, line in enumerate(chunks[start:start+13]):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                paragraph.text = line; paragraph.font.size = Pt(17)
        deck.save(out)
    return out.getvalue()
