import io
import zipfile
import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from app.core.errors import AppError
from app.services.document_service import extract, validate_file
from app.services.retrieval_service import retrieve


def make_file(ext):
    output = io.BytesIO()
    if ext == '.docx':
        d = Document(); d.add_paragraph('Invoice processing workflow'); d.add_table(rows=1, cols=1).cell(0,0).text = 'Accounting'; d.save(output)
    elif ext == '.xlsx':
        w = Workbook(); w.active.append(['Invoice', 'Owner']); w.active.append(['INV-01', 'Accounting']); w.save(output)
    elif ext == '.pdf':
        w = PdfWriter(); page = w.add_blank_page(300, 300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        content = DecodedStreamObject(); content.set_data(b'BT /F1 12 Tf 20 200 Td (Invoice processing workflow) Tj ET')
        page[NameObject('/Contents')] = w._add_object(content); w.write(output)
    else:
        return b'Invoice,Owner\nINV-01,Accounting' if ext == '.csv' else b'Invoice processing workflow'
    return output.getvalue()


@pytest.mark.parametrize('ext', ['.pdf', '.docx', '.txt', '.csv', '.xlsx'])
def test_supported_formats_extract_and_cite(ext):
    data = make_file(ext)
    validate_file('workflow' + ext, data)
    chunks = extract(data, ext)
    assert 'Invoice' in chunks[0]['text']
    assert chunks[0]['metadata']['page']


def test_scanned_pdf_gives_actionable_error():
    out = io.BytesIO(); w = PdfWriter(); w.add_blank_page(300, 300); w.write(out)
    with pytest.raises(AppError, match='OCR'):
        extract(out.getvalue(), '.pdf')


def test_office_type_mismatch_rejected():
    with pytest.raises(AppError):
        validate_file('wrong.xlsx', make_file('.docx'))


def test_retrieval_is_project_scoped(setup):
    _, store, ai, _ = setup
    store.db.document_chunks.insert_many([{'project_id': 'a', 'text': 'Invoice processing workflow'}, {'project_id': 'b', 'text': 'Private invoice information'}])
    results, warnings = retrieve(store, ai, 'a', 'Invoice')
    assert len(results) == 1
    assert results[0]['project_id'] == 'a'
    assert not warnings
