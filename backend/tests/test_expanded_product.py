import io
import json
import zipfile
from datetime import timedelta
from xml.etree import ElementTree
import pytest
from bson import ObjectId
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from app.core.auth import signer
from app.core.errors import AppError
from app.core.passwords import verify_password
from app.models.deliverables import KINDS, Deliverable
from app.repositories.store import now
from tests.fakes import deliverable_fixture


def register(client, email='owner@example.com', name='Test Owner'):
    client.headers['Origin']='http://localhost:3000'
    response=client.post('/api/auth/signup',json={'name':name,'email':email,'password':'test-only-password-123'})
    assert response.status_code==201, response.text
    return response.json(), client.cookies.get('shift_session')


def ready(client,pid):
    assert client.post(f'/api/projects/{pid}/discovery/next').status_code==200
    assert client.post(f'/api/projects/{pid}/analysis/run').status_code==202
    assert client.get(f'/api/projects/{pid}').json()['status']=='BLUEPRINT_READY'


def test_real_email_flow_hashing_sessions_isolation(setup,project):
    client,store,_,settings=setup
    account, token=register(client,' OWNER@EXAMPLE.COM ')
    assert account['email']=='owner@example.com' and not account['local']
    assert 'HttpOnly' in client.post('/api/auth/login',json={'email':account['email'],'password':'test-only-password-123'}).headers['set-cookie']
    record=store.db.users.find_one({'email':account['email']})
    assert verify_password('test-only-password-123',record['password_hash'])
    assert 'password' not in json.dumps(account)
    assert client.get('/api/projects').json()==[]
    assert client.get(f'/api/projects/{project}').status_code==404
    assert client.get('/api/auth/me').json()['id']==account['id']
    assert client.post('/api/auth/signup',json={'name':'Duplicate','email':'owner@example.com','password':'test-only-password-123'}).status_code==409
    settings.allow_local_access=False
    assert client.post('/api/auth/logout').status_code==200
    assert client.get('/api/auth/me').status_code==401
    assert client.post('/api/auth/login',json={'email':'owner@example.com','password':'wrong'}).status_code==401
    assert client.post('/api/auth/login',json={'email':'owner@example.com','password':'test-only-password-123'}).status_code==200
    store.db.users.update_one({'_id':record['_id']},{'$inc':{'session_version':1}})
    assert client.get('/api/auth/me').status_code==401


def test_auth_validation_rate_limit_and_cross_origin(setup):
    client,*_=setup
    for body in [{'name':'A','email':'x@example.com','password':'test-only-password-123'}, {'name':'Valid','email':'invalid','password':'test-only-password-123'}, {'name':'Valid','email':'x@example.com','password':'short'}]:
        assert client.post('/api/auth/signup',json=body).status_code==422
    assert client.post('/api/auth/login',json={'email':'x@example.com','password':'wrong'},headers={'Origin':'https://attacker.invalid'}).status_code==403
    for _ in range(10):
        assert client.post('/api/auth/login',json={'email':'absent@example.com','password':'wrong'}).status_code==401
    assert client.post('/api/auth/login',json={'email':'absent@example.com','password':'wrong'}).status_code==429


@pytest.mark.parametrize('role', ['editor','reviewer','viewer','admin'])
def test_team_permissions_invitations_notifications_and_revocation(setup,role):
    client,store,_,_=setup
    owner,owner_cookie=register(client)
    wid=client.post('/api/workspaces',json={'name':'Pilot team','organization':'Example'}).json()['id']
    pid=client.post('/api/projects',json={'name':'Team invoice','initial_problem':'Review manual invoice imports safely.','workspace_id':wid}).json()['id']
    content=deliverable_fixture()
    assert client.post(f'/api/projects/{pid}/deliverables/business',json={'base_version':0,'content':content}).status_code==200
    invite=client.post(f'/api/workspaces/{wid}/invitations',json={'role':role}).json()['token']
    assert invite not in str(store.db.invitations.find_one())
    other,other_cookie=register(client,'member@example.com','Team Member')
    assert client.get(f'/api/projects/{pid}').status_code==404
    assert client.post('/api/invitations/accept',json={'token':invite}).status_code==200
    assert client.post('/api/invitations/accept',json={'token':invite}).status_code==400
    assert client.get(f'/api/projects/{pid}').json()['access_role']==role
    assert len(client.get('/api/projects').json())==1
    edited=client.post(f'/api/projects/{pid}/deliverables/business',json={'base_version':1,'content':content})
    assert edited.status_code==(200 if role in ['editor','admin'] else 403)
    version=2 if edited.status_code==200 else 1
    reviewed=client.post(f'/api/projects/{pid}/deliverables/business/review',json={'version':version,'decision':'approved','note':'Validated assumptions in pilot.'})
    assert reviewed.status_code==(200 if role in ['reviewer','admin'] else 403)
    assert client.get(f'/api/workspaces/{wid}/admin').status_code==(200 if role=='admin' else 403)
    assert client.post(f'/api/workspaces/{wid}/members/{owner["id"]}',json={'role':'viewer'}).status_code==403
    assert client.post(f'/api/projects/{pid}/comments',json={'content':'Confirm accounting import access.'}).status_code==201
    client.cookies.clear()
    client.cookies.set('shift_session',owner_cookie)
    assert client.get('/api/notifications').json()
    assert client.delete(f'/api/workspaces/{wid}/members/{other["id"]}').status_code==200
    client.cookies.clear()
    client.cookies.set('shift_session',other_cookie)
    assert client.get(f'/api/projects/{pid}').status_code==404
    assert client.get('/api/projects').json()==[]


@pytest.mark.parametrize('kind', KINDS)
def test_generation_version_review_exports_and_stale_context(setup,project,kind):
    client,store,ai,_=setup
    root=f'/api/projects/{project}/deliverables/{kind}'
    assert client.post(root+'/generate',json={}).status_code==409
    ready(client,project)
    assert client.post(root+'/generate',json={'language':'hi','instructions':'Keep human approval.'}).status_code==202
    state=client.get(f'/api/projects/{project}/deliverables').json()
    assert state['jobs'][kind]['status']=='complete', state
    artifact=next(i['artifact'] for i in state['items'] if i['kind']==kind)
    assert artifact['version']==1 and artifact['language']=='hi' and artifact['ai_reviews']
    assert client.post(root+'/review',json={'version':1,'decision':'approved','note':'Checked source evidence.'}).status_code==200
    content=artifact['content'];content['summary']='Changed after stakeholder feedback.'
    assert client.post(root,json={'base_version':1,'content':content}).json()['version']==2
    assert client.post(root,json={'base_version':1,'content':content}).status_code==409
    assert len(client.get(root+'/versions').json())==2
    assert client.post(root+'/review',json={'version':1,'decision':'approved','note':'Old review'}).status_code==409
    assert client.post(f'/api/projects/{project}/chat',json={'content':'The accounting system changed; revalidate API access.'}).status_code==200
    assert next(i['stale'] for i in client.get(f'/api/projects/{project}/deliverables').json()['items'] if i['kind']==kind)
    assert client.post(root+'/review',json={'version':2,'decision':'approved','note':'Stale review'}).status_code==409
    assert client.get(f'/api/projects/{project}/export/{kind}/json?version=1').json()['summary']!='Changed after stakeholder feedback.'


def test_pack_red_team_bound_and_provider_failure_preserve_versions(setup,project,monkeypatch):
    client,store,ai,_=setup
    ready(client,project); ai.always_revise=True
    assert client.post(f'/api/projects/{project}/deliverables/all/generate',json={}).status_code==202
    assert store.db.artifacts.count_documents({'project_id':project})==7
    assert all(len(a['ai_reviews'])==3 for a in store.db.artifacts.find({'project_id':project}))
    assert all(a['ai_reviews'][-1]['findings'][0]['requires_revision'] for a in store.db.artifacts.find({'project_id':project}))
    monkeypatch.setattr(ai,'generate_structured',lambda *a,**k: (_ for _ in ()).throw(AppError('Provider unavailable',503)))
    assert client.post(f'/api/projects/{project}/deliverables/business/generate',json={}).status_code==202
    state=client.get(f'/api/projects/{project}/deliverables').json()
    assert not state['busy'] and state['jobs']['business']['status']=='error'
    assert store.db.artifacts.count_documents({'project_id':project,'kind':'business'})==1


def test_office_bpmn_zip_exports_round_trip_and_formula_safety(setup,project):
    client,*_=setup
    data=deliverable_fixture();data['sections'][0]['tables'][0]['rows'][0][0]='=HYPERLINK("https://invalid")'
    client.post(f'/api/projects/{project}/deliverables/process',json={'base_version':0,'content':data})
    def get(fmt):
        response=client.get(f'/api/projects/{project}/export/process/{fmt}')
        assert response.status_code==200,response.text[:200]
        return io.BytesIO(response.content)
    assert any('Current' in p.text for p in Document(get('docx')).paragraphs)
    wb=load_workbook(get('xlsx'))
    assert wb['Table 1']['A1'].value=='Work item'
    assert wb['Table 1']['A2'].value.startswith('=HYPERLINK') and wb['Table 1']['A2'].data_type=='s'
    assert len(Presentation(get('pptx')).slides)>0
    xml=ElementTree.parse(get('bpmn'));ns={'b':'http://www.omg.org/spec/BPMN/20100524/MODEL'}
    assert len(xml.findall('.//b:sequenceFlow',ns))==2
    with zipfile.ZipFile(get('zip')) as archive:
        assert {'schema.sql','openapi.yaml','deliverable.json','report.md'}.issubset(archive.namelist())
        assert archive.testzip() is None


def test_pptx_ingestion(setup,project):
    client,*_=setup
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Invoice process';slide.placeholders[1].text='Operations reviews invoice records before importing them.'
    payload=io.BytesIO();deck.save(payload)
    response=client.post(f'/api/projects/{project}/documents',files={'file':('workflow.pptx',payload.getvalue())})
    assert response.status_code==202,response.text
    assert client.get(f'/api/projects/{project}/documents').json()[0]['status']=='processed'


def test_backup_restore_preserves_original_and_does_not_import_approval(setup):
    client,store,_,_=setup
    wid=client.post('/api/workspaces',json={'name':'Backup team'}).json()['id']
    pid=client.post('/api/projects',json={'name':'Backup project','initial_problem':'Manual invoice processing workflow.','workspace_id':wid}).json()['id']
    client.post(f'/api/projects/{pid}/documents',files={'file':('workflow.txt',b'Operations reviews records before importing invoice data.')})
    ready(client,pid)
    client.post(f'/api/projects/{pid}/deliverables/business/generate',json={})
    client.post(f'/api/projects/{pid}/deliverables/business/review',json={'version':1,'decision':'approved','note':'Reviewed.'})
    backup=client.get(f'/api/workspaces/{wid}/backup')
    assert backup.status_code==200
    assert 'password_hash' not in backup.text and 'encrypted_token' not in backup.text
    response=client.post(f'/api/workspaces/{wid}/restore',json=backup.json())
    assert response.status_code==201,response.text
    new=response.json()['project_ids'][0];assert new!=pid
    assert client.get(f'/api/projects/{pid}').json()['status']=='BLUEPRINT_READY'
    assert client.get(f'/api/projects/{new}').json()['status']=='DISCOVERY'
    assert client.get(f'/api/projects/{new}/collaboration').json()['reviews']==[]
    assert client.get(f'/api/projects/{new}/restored-history').json()['reviews_archive']
    assert client.get(f'/api/projects/{new}/deliverables').json()['items'][0]['stale']
    assert store.db.document_chunks.count_documents({'project_id':new})>0
    bad=backup.json();bad['projects'][0]['document_chunks'][0]['document_id']='missing'
    before=store.db.projects.count_documents({})
    assert client.post(f'/api/workspaces/{wid}/restore',json=bad).status_code==422
    assert store.db.projects.count_documents({})==before


def test_localization_allowlist_cache_and_outcomes(setup,project):
    client,store,ai,_=setup
    before=len(ai.calls)
    assert client.post('/api/localization/hi',json={'texts':['secret project content']}).status_code==400
    assert len(ai.calls)==before
    first=client.post('/api/localization/hi',json={'texts':['Projects']})
    assert first.status_code==200,first.text
    assert client.post('/api/localization/hi',json={'texts':['Projects']}).json()==first.json()
    assert ai.calls.count('TranslationOutput')==1
    assert client.post(f'/api/projects/{project}/outcomes',json={'metric':'Entry time','unit':'minutes','baseline':30,'target':10,'observed':15}).status_code==201
    assert client.get('/api/transformation').json()[0]['outcomes'][0]['observed']==15


def test_microsoft_connection_encrypted_import_and_disconnect(setup,monkeypatch):
    from app.api import integrations
    client,store,_,_=setup
    wid=client.post('/api/workspaces',json={'name':'Connected team'}).json()['id']
    pid=client.post('/api/projects',json={'name':'Planner pilot','initial_problem':'Coordinate assigned invoice tasks.','workspace_id':wid}).json()['id']
    token='test-only-microsoft-token-12345'
    calls=[]
    def graph(value,path):
        assert value==token;calls.append(path)
        return {'displayName':'Example'} if path.startswith('me?') else {'value':[{'id':'t1','title':'Review invoice','percentComplete':0}]}
    monkeypatch.setattr(integrations,'graph',graph)
    assert client.post(f'/api/workspaces/{wid}/integrations/microsoft_graph',json={'access_token':token}).status_code==200
    assert token not in str(store.db.integrations.find_one())
    assert token not in client.get(f'/api/workspaces/{wid}/integrations').text
    assert client.post(f'/api/projects/{pid}/integrations/microsoft_graph/import').status_code==202
    assert client.get(f'/api/projects/{pid}/documents').json()[0]['status']=='processed'
    assert calls==['me?$select=displayName','me/planner/tasks']
    assert client.delete(f'/api/workspaces/{wid}/integrations/microsoft_graph').status_code==200
    assert not client.get(f'/api/workspaces/{wid}/integrations').json()['microsoft_graph']['connected']
